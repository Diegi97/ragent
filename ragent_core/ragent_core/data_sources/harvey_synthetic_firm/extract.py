"""Text extraction for the four file formats in Harvey's published firm corpus."""

from datetime import date, datetime, time
from email import policy
from email.header import decode_header, make_header
from email.parser import BytesParser
from pathlib import Path
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SUPPORTED_EXTENSIONS = {".docx", ".eml", ".xlsx", ".pptx"}


def _table_rows(table) -> str:
    return "\n".join(
        " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
        for row in table.rows
    )


def extract_docx(path: Path) -> str:
    document = Document(path)
    blocks = []
    # iter_inner_content preserves the placement of tables between paragraphs.
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            blocks.append(_table_rows(block))
        else:
            blocks.append(block.text)
    # These parts are not exposed through Document.paragraphs. Read their text,
    # but never styles, document properties, comments, or deleted revision text.
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for part in document.part.package.parts:
        name = Path(str(part.partname)).name
        if name.startswith(
            ("header", "footer", "footnotes", "endnotes")
        ) and name.endswith(".xml"):
            root = ElementTree.fromstring(part.blob)
            paragraphs = [
                "".join(node.text or "" for node in paragraph.iter(f"{namespace}t"))
                for paragraph in root.iter(f"{namespace}p")
            ]
            text = "\n\n".join(p for p in paragraphs if p.strip())
            if text:
                blocks.append(f"## {name}\n\n{text}")
    return "\n\n".join(block for block in blocks if block.strip())


def _cell_value(value) -> str:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value).replace("\n", " / ")


def extract_xlsx(path: Path) -> str:
    formulas = load_workbook(path, read_only=True, data_only=False)
    cached = None
    try:
        cached = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for sheet in formulas:
            rows = []
            values = cached[sheet.title].iter_rows()
            for row, cached_row in zip(sheet.iter_rows(), values, strict=True):
                cells = []
                for cell, cached_cell in zip(row, cached_row, strict=True):
                    if cell.value is None:
                        continue
                    text = _cell_value(cell.value)
                    if cell.data_type == "f":
                        result = cached_cell.value
                        text += (
                            f" [cached value: {_cell_value(result)}]"
                            if result is not None
                            else " [no cached value]"
                        )
                    elif cell.number_format != "General" and isinstance(
                        cell.value, (int, float)
                    ):
                        text += f" [format: {cell.number_format}]"
                    cells.append(f"{cell.coordinate}: {text}")
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"## Sheet: {sheet.title}\n\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def _shape_text(shapes) -> list[str]:
    blocks = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            blocks.extend(_shape_text(shape.shapes))
        elif shape.has_table:
            blocks.append(_table_rows(shape.table))
        elif shape.has_text_frame:
            blocks.append(shape.text)
    return [block for block in blocks if block.strip()]


def extract_pptx(path: Path) -> str:
    presentation = Presentation(path)
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks = _shape_text(slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                blocks.append(f"### Speaker notes\n\n{frame.text}")
        if blocks:
            slides.append(f"## Slide {index}\n\n" + "\n\n".join(blocks))
    return "\n\n".join(slides)


def extract_eml(path: Path) -> str:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    headers = []
    for key in (
        "From",
        "To",
        "Cc",
        "Bcc",
        "Date",
        "Subject",
        "Reply-To",
        "Message-ID",
        "In-Reply-To",
        "References",
    ):
        # Some released emails separate recipients with semicolons. HeaderRegistry
        # can silently drop those addresses; decode the original headers instead.
        values = [
            " ".join(
                str(
                    make_header(
                        decode_header(
                            value.encode("utf-8", errors="surrogateescape").decode(
                                "utf-8", errors="replace"
                            )
                        )
                    )
                ).splitlines()
            )
            for name, value in message.raw_items()
            if name.casefold() == key.casefold()
        ]
        if values:
            headers.append(f"{key}: " + "; ".join(values))
    body = message.get_body(preferencelist=("plain", "html"))
    if body is None:
        raise ValueError(f"Email has no text body: {path}")
    text = body.get_content()
    if body.get_content_type() == "text/html":
        soup = BeautifulSoup(text, "html.parser")
        for node in soup(["script", "style"]):
            node.decompose()
        text = soup.get_text(separator="\n", strip=True)
    attachments = [
        part.get_filename()
        for part in message.iter_attachments()
        if part.get_filename()
    ]
    if attachments:
        headers.append("Attachments (names only): " + "; ".join(attachments))
    if not text.strip():
        raise ValueError(f"Email has an empty text body: {path}")
    return "\n".join(headers) + "\n\n" + text.strip()


def extract_document(path: Path) -> str:
    readers = {
        ".docx": extract_docx,
        ".xlsx": extract_xlsx,
        ".pptx": extract_pptx,
        ".eml": extract_eml,
    }
    text = readers[path.suffix.lower()](path).strip()
    if not text:
        raise ValueError(f"No text extracted from {path}")
    return text
